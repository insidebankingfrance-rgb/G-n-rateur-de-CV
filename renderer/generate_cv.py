"""Render a normalized CV JSON to an Inside Circle .pptx (FR + EN slides).

Usage:
    python renderer/generate_cv.py <input.json> [--out outputs/<name>.pptx]
    python renderer/generate_cv.py data/ --out outputs/    # batch mode

Garanties contrôlées en sortie :
- Toute typo de contenu est ≥ 10pt (`FS_FLOOR`). Lève AssertionError sinon.
- Le contenu est plafonné via `MAX_*` dans `theme.py` pour rentrer en 1 page.
  Les éléments tronqués sont loggués (stderr).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Iterable

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

from theme import (
    CYAN_ACCENT,
    FONT_FAMILY,
    FOOTER_H,
    FOOTER_TEXT,
    FOOTER_TEXT_EN,
    FS_BODY,
    FS_BODY_SMALL,
    FS_DOMAIN,
    FS_FLOOR,
    FS_FOOTER,
    FS_INITIALS,
    FS_SECTION,
    FS_TITLE,
    HEADER_H,
    LOGO_FILENAME,
    LOGO_H,
    LOGO_RESERVE_W,
    LOGO_RIGHT_PAD,
    LOGO_TOP,
    MARGIN,
    MAX_BULLETS_PER_EXP,
    MAX_ENGAGEMENTS,
    MAX_EXPERIENCES,
    MAX_EXPERTISE,
    MAX_HOBBIES,
    NAVY_SIDEBAR,
    SEPARATOR_W,
    SIDEBAR_W,
    SLIDE_H,
    SLIDE_W,
    WHITE,
    WHITE_SOFT,
)
from i18n import LABELS


_warnings: list[str] = []


def _warn(msg: str) -> None:
    _warnings.append(msg)
    print(f"  ⚠ {msg}", file=sys.stderr)


def _check_min_font(size, where: str) -> None:
    if size is None:
        return
    if size < FS_FLOOR:
        raise AssertionError(
            f"Font size {size.pt}pt below floor {FS_FLOOR.pt}pt at {where}"
        )


# ─── Low-level helpers ───────────────────────────────────────────────────────
def _set_fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _add_rect(slide, left, top, width, height, color: RGBColor):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    _set_fill(shp, color)
    shp.shadow.inherit = False
    return shp


def _add_textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.paragraphs[0].text = ""
    return tb, tf


def _run(paragraph, text: str, *, bold=False, italic=False,
         size=FS_BODY, color=WHITE, font=FONT_FAMILY, allow_small=False):
    if not allow_small:
        _check_min_font(size, f'run "{text[:40]}"')
    r = paragraph.add_run()
    r.text = text
    f = r.font
    f.name = font
    f.size = size
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    return r


def _new_para(tf, *, align=PP_ALIGN.LEFT, space_after=Pt(2), level=0):
    if not tf.paragraphs[0].runs and tf.paragraphs[0].text == "":
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = align
    p.space_after = space_after
    p.level = level
    return p


def _section_title(tf, label: str, *, first=False):
    p = _new_para(tf, space_after=Pt(3))
    if not first:
        p.space_before = Pt(6)
    _run(p, label, bold=True, size=FS_SECTION, color=CYAN_ACCENT)


def _bullet(tf, text: str, *, size=FS_BODY, color=WHITE, level=0,
            bold=False, italic=False):
    p = _new_para(tf, space_after=Pt(1), level=level)
    _run(p, "•  ", bold=True, size=size, color=CYAN_ACCENT)
    _run(p, text, size=size, color=color, bold=bold, italic=italic)


# ─── Anonymization / truncation ──────────────────────────────────────────────
def initials_from_name(first: str, last: str) -> str:
    f, l = (first or "").strip(), (last or "").strip()
    if not f and not l:
        return "?.?."
    return f"{f[:1].upper()}.{l[:1].upper()}."


def _truncate_payload(payload: dict, who: str) -> dict:
    """Apply MAX_* caps by dropping WHOLE items only — jamais de troncature
    mid-string avec "…". Le texte des bullets / summary / descriptions doit
    être écrit à la bonne longueur dès la phase d'extraction."""
    out = dict(payload)

    exp = out.get("experiences", [])
    if len(exp) > MAX_EXPERIENCES:
        dropped = [e.get("employer", "?") for e in exp[MAX_EXPERIENCES:]]
        _warn(f"{who}: {len(exp) - MAX_EXPERIENCES} experience(s) dropped "
              f"(oldest first): {', '.join(dropped)}")
        exp = exp[:MAX_EXPERIENCES]
    exp = [dict(e) for e in exp]
    for e in exp:
        ach = e.get("achievements", [])
        if len(ach) > MAX_BULLETS_PER_EXP:
            _warn(f"{who}: {e.get('employer', '?')} — "
                  f"{len(ach) - MAX_BULLETS_PER_EXP} bullet(s) dropped")
            e["achievements"] = ach[:MAX_BULLETS_PER_EXP]
    out["experiences"] = exp

    for key, cap in (("expertise", MAX_EXPERTISE),
                     ("hobbies", MAX_HOBBIES),
                     ("engagements", MAX_ENGAGEMENTS)):
        if key in out and len(out[key]) > cap:
            _warn(f"{who}: {key} capped to {cap} (was {len(out[key])})")
            out[key] = out[key][:cap]
    return out


# ─── Overflow estimator ──────────────────────────────────────────────────────
# Each column has a vertical budget (in pt). We estimate the rendered height
# from font sizes and wrap, and warn if we exceed the budget.

# Char widths derived empirically for ~10pt Alegreya Sans / sans fallbacks.
# CPL recalibrated on 12pt Alegreya Sans / sans-serif. Sidebar 35.2% × 13.33"
# = 4.7" → ~42 chars/line. Main 64.8% × 13.33" = 8.6" → ~80 chars/line.
_CPL_SIDEBAR_BODY  = 42
_CPL_MAIN_BODY     = 80
_LINE_PT           = 14   # 12pt + ~17% leading
_SECTION_PT        = 19
_PARA_GAP_PT       = 1

# EMU per pt = 12700
_AVAIL_PT_MAIN    = (SLIDE_H - HEADER_H - FOOTER_H - 160000) / 12700
_AVAIL_PT_SIDEBAR = _AVAIL_PT_MAIN - 40  # initials block (~0.5") eats the top


def _wrapped_lines(text: str, cpl: int) -> int:
    if not text:
        return 0
    # naive: count chars / cpl, min 1
    return max(1, (len(text) + cpl - 1) // cpl)


def _estimate_sidebar_pt(payload: dict) -> float:
    pt = 0.0
    if payload.get("expertise"):
        pt += _SECTION_PT + len(payload["expertise"]) * _LINE_PT
    if payload.get("languages"):
        pt += _SECTION_PT + len(payload["languages"]) * _LINE_PT
    if payload.get("education"):
        pt += _SECTION_PT
        for e in payload["education"]:
            year = (e.get("year", "") if isinstance(e, dict) else "")
            school = (e.get("school", "") if isinstance(e, dict) else str(e))
            degree = (e.get("degree", "") if isinstance(e, dict) else "")
            pt += _wrapped_lines(f"{year} {school}", _CPL_SIDEBAR_BODY) * _LINE_PT
            if degree:
                pt += _wrapped_lines(degree, _CPL_SIDEBAR_BODY) * _LINE_PT
    if payload.get("hobbies"):
        pt += _SECTION_PT + sum(
            _wrapped_lines(h, _CPL_SIDEBAR_BODY) for h in payload["hobbies"]
        ) * _LINE_PT
    if payload.get("engagements"):
        pt += _SECTION_PT
        for e in payload["engagements"]:
            if isinstance(e, dict):
                line = f"{e.get('title','')} {e.get('description','')}"
            else:
                line = str(e)
            pt += _wrapped_lines(line, _CPL_SIDEBAR_BODY) * _LINE_PT + _PARA_GAP_PT
    return pt


def _estimate_main_pt(payload: dict) -> float:
    pt = 0.0
    if (payload.get("summary") or "").strip():
        pt += _SECTION_PT + _wrapped_lines(payload["summary"], _CPL_MAIN_BODY) * _LINE_PT
    if payload.get("experiences"):
        pt += _SECTION_PT
        for exp in payload["experiences"]:
            header = (f"{exp.get('employer','')} {exp.get('role','')} "
                      f"{exp.get('duration','')}")
            pt += _wrapped_lines(header, _CPL_MAIN_BODY) * _LINE_PT + _PARA_GAP_PT
            for b in exp.get("achievements", []):
                pt += _wrapped_lines(b, _CPL_MAIN_BODY - 4) * _LINE_PT + _PARA_GAP_PT
    if payload.get("references"):
        ref = ", ".join(payload["references"])
        pt += _SECTION_PT + _wrapped_lines(ref, _CPL_MAIN_BODY) * _LINE_PT
    return pt


def _check_overflow(payload: dict, who: str) -> None:
    main_pt = _estimate_main_pt(payload)
    side_pt = _estimate_sidebar_pt(payload)
    if main_pt > _AVAIL_PT_MAIN:
        _warn(f"{who}: main column estimated {main_pt:.0f}pt > "
              f"budget {_AVAIL_PT_MAIN:.0f}pt — risque de débordement")
    if side_pt > _AVAIL_PT_SIDEBAR:
        _warn(f"{who}: sidebar estimated {side_pt:.0f}pt > "
              f"budget {_AVAIL_PT_SIDEBAR:.0f}pt — risque de débordement")


def _adaptive_main_trim(payload: dict, who: str) -> dict:
    """When main column overflow is predicted, drop content by descending
    priority. Only references get dropped (bullets/experiences are protected
    — the truncation in _truncate_payload already applied).
    Le seuil est plus strict que _AVAIL_PT_MAIN pour laisser une marge de
    sécurité (l'estimateur peut sous-estimer de quelques pt)."""
    payload = dict(payload)
    safety = _AVAIL_PT_MAIN - 20
    if _estimate_main_pt(payload) <= safety:
        return payload
    if payload.get("references"):
        _warn(f"{who}: references dropped to keep main in 1 page")
        payload["references"] = []
    return payload


def _adaptive_sidebar_trim(payload: dict, who: str) -> dict:
    """When sidebar overflow is predicted, drop content by descending priority:
    hobbies first, then engagement descriptions, then engagements entirely.
    Expertise / languages / education are kept untouched."""
    payload = dict(payload)
    if _estimate_sidebar_pt(payload) <= _AVAIL_PT_SIDEBAR:
        return payload
    if payload.get("hobbies"):
        _warn(f"{who}: hobbies dropped to keep sidebar in 1 page")
        payload["hobbies"] = []
    if _estimate_sidebar_pt(payload) <= _AVAIL_PT_SIDEBAR:
        return payload
    if payload.get("engagements"):
        new = []
        for e in payload["engagements"]:
            if isinstance(e, dict):
                e = {**e, "description": ""}
            new.append(e)
        payload["engagements"] = new
        _warn(f"{who}: engagement descriptions dropped to keep sidebar in 1 page")
    if _estimate_sidebar_pt(payload) <= _AVAIL_PT_SIDEBAR:
        return payload
    if payload.get("engagements"):
        payload["engagements"] = payload["engagements"][:2]
        _warn(f"{who}: engagements capped to 2 to keep sidebar in 1 page")
    return payload


# ─── Slide builders ──────────────────────────────────────────────────────────
def _build_background(slide):
    """Fond uniforme NAVY_SIDEBAR sur tout le slide."""
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY_SIDEBAR)


def _build_separator(slide):
    """Trait cyan vertical entre la sidebar et la main column."""
    _add_rect(slide, SIDEBAR_W, 0, SEPARATOR_W, SLIDE_H, CYAN_ACCENT)


def _build_footer_mask(slide):
    """Bandeau navy plein-largeur au-dessus du footer pour masquer
    tout débordement de textbox dans la zone du bas."""
    _add_rect(slide, 0, SLIDE_H - FOOTER_H, SLIDE_W, FOOTER_H, NAVY_SIDEBAR)


def _build_logo(slide, logo_path: Path | None):
    """Place the Inside Circle logo in the top-right corner of the slide.
    Silently no-op if the file is missing. The image is resized to LOGO_H
    keeping its native aspect ratio (so logos of any aspect work)."""
    if logo_path is None or not logo_path.is_file():
        return
    try:
        from PIL import Image
        with Image.open(logo_path) as im:
            w, h = im.size
    except Exception:
        # No PIL or unreadable — fall back to assumed square
        w = h = 1
    aspect = w / h if h else 1.0
    logo_w = int(LOGO_H * aspect)
    left = SLIDE_W - LOGO_RIGHT_PAD - logo_w
    slide.shapes.add_picture(str(logo_path), left, LOGO_TOP,
                             width=logo_w, height=LOGO_H)


def _resolve_logo() -> Path | None:
    """Locate assets/inside_circle_logo.png relative to the repo root."""
    here = Path(__file__).resolve().parent
    candidate = here.parent / "assets" / LOGO_FILENAME
    return candidate if candidate.is_file() else None


def _build_header(slide, initials: str, title: str, domain: str):
    # Initiales dans la sidebar
    tb, tf = _add_textbox(
        slide,
        MARGIN, Emu(260000),
        SIDEBAR_W - 2 * MARGIN, Emu(800000),
    )
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    _run(p, initials, bold=True, size=FS_INITIALS, color=CYAN_ACCENT)

    # Titre + domaine — width réduit pour laisser la place au logo top-right
    tb, tf = _add_textbox(
        slide,
        SIDEBAR_W + MARGIN, Emu(300000),
        SLIDE_W - SIDEBAR_W - MARGIN - LOGO_RESERVE_W, Emu(800000),
    )
    p = tf.paragraphs[0]
    _run(p, title or "", bold=True, size=FS_TITLE, color=WHITE)
    if domain:
        p2 = tf.add_paragraph()
        p2.space_before = Pt(2)
        _run(p2, domain, size=FS_DOMAIN, color=CYAN_ACCENT, italic=True)


def _build_footer(slide, lang: str):
    text = FOOTER_TEXT if lang == "fr" else FOOTER_TEXT_EN
    tb, tf = _add_textbox(
        slide,
        MARGIN, SLIDE_H - FOOTER_H,
        SLIDE_W - 2 * MARGIN, FOOTER_H,
    )
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    # FS_FOOTER (9pt) — exception assumée pour la mention légale.
    _run(p, text, italic=True, size=FS_FOOTER, color=WHITE_SOFT,
         allow_small=True)


def _build_sidebar_content(slide, cv: dict, lang: str):
    L = LABELS[lang]
    top = HEADER_H + Emu(80000)
    height = SLIDE_H - top - FOOTER_H - Emu(80000)
    tb, tf = _add_textbox(
        slide, MARGIN, top, SIDEBAR_W - 2 * MARGIN, height,
    )
    tf.word_wrap = True
    first = True

    def section(label, render):
        nonlocal first
        _section_title(tf, label, first=first)
        first = False
        render()

    if cv.get("expertise"):
        def _r():
            for e in cv["expertise"]:
                _bullet(tf, e, size=FS_BODY)
        section(L["expertise"], _r)

    if cv.get("languages"):
        def _r():
            for lng in cv["languages"]:
                if isinstance(lng, dict):
                    name, level = lng.get("name", ""), lng.get("level", "")
                    txt = f"{name} — {level}" if level else name
                else:
                    txt = str(lng)
                _bullet(tf, txt, size=FS_BODY)
        section(L["languages"], _r)

    if cv.get("education"):
        def _r():
            for it in cv["education"]:
                if isinstance(it, dict):
                    year, school, degree = (it.get("year", ""),
                                            it.get("school", ""),
                                            it.get("degree", ""))
                    p = _new_para(tf, space_after=Pt(1))
                    if year:
                        _run(p, f"{year}  ", bold=True, size=FS_BODY,
                             color=CYAN_ACCENT)
                    if school:
                        _run(p, school, bold=True, size=FS_BODY, color=WHITE)
                    if degree:
                        p2 = _new_para(tf, space_after=Pt(3))
                        _run(p2, degree, size=FS_BODY, color=WHITE_SOFT,
                             italic=True)
                else:
                    _bullet(tf, str(it), size=FS_BODY)
        section(L["education"], _r)

    if cv.get("hobbies"):
        def _r():
            for h in cv["hobbies"]:
                _bullet(tf, h, size=FS_BODY)
        section(L["hobbies"], _r)

    if cv.get("engagements"):
        def _r():
            for e in cv["engagements"]:
                if isinstance(e, dict):
                    title, desc = e.get("title", ""), e.get("description", "")
                    p = _new_para(tf, space_after=Pt(2))
                    _run(p, "• ", bold=True, size=FS_BODY, color=CYAN_ACCENT)
                    if title:
                        _run(p, title, bold=True, size=FS_BODY, color=WHITE)
                    if desc:
                        _run(p, f" — {desc}", size=FS_BODY, color=WHITE_SOFT,
                             italic=True)
                else:
                    _bullet(tf, str(e), size=FS_BODY)
        section(L["engagements"], _r)


def _build_main_content(slide, cv: dict, lang: str):
    L = LABELS[lang]
    left = SIDEBAR_W + MARGIN
    top = HEADER_H + Emu(80000)
    width = SLIDE_W - SIDEBAR_W - 2 * MARGIN
    height = SLIDE_H - top - FOOTER_H - Emu(80000)
    tb, tf = _add_textbox(slide, left, top, width, height)
    tf.word_wrap = True
    first = True

    def section(label, render):
        nonlocal first
        _section_title(tf, label, first=first)
        first = False
        render()

    summary = (cv.get("summary") or "").strip()
    if summary:
        def _r():
            p = _new_para(tf, space_after=Pt(3))
            _run(p, summary, size=FS_BODY, color=WHITE)
        section(L["summary"], _r)

    if cv.get("experiences"):
        def _r():
            for exp in cv["experiences"]:
                employer, role, duration = (exp.get("employer", ""),
                                            exp.get("role", ""),
                                            exp.get("duration", ""))
                p = _new_para(tf, space_after=Pt(1))
                _run(p, "▸ ", bold=True, size=FS_BODY, color=CYAN_ACCENT)
                if employer:
                    _run(p, employer, bold=True, size=FS_BODY, color=WHITE)
                if role:
                    _run(p, f" — {role}", size=FS_BODY, color=WHITE)
                if duration:
                    _run(p, f"  ({duration})", italic=True,
                         size=FS_BODY_SMALL, color=CYAN_ACCENT)
                for bul in exp.get("achievements", []):
                    _bullet(tf, bul, size=FS_BODY_SMALL, level=1)
        section(L["experience"], _r)

    if cv.get("references"):
        def _r():
            p = _new_para(tf, space_after=Pt(2))
            _run(p, ", ".join(cv["references"]), size=FS_BODY,
                 color=WHITE, italic=True)
        section(L["references"], _r)


# ─── Top-level render ────────────────────────────────────────────────────────
def _build_slide(prs: Presentation, cv: dict, lang: str, logo: Path | None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _build_background(slide)
    _build_separator(slide)
    _build_logo(slide, logo)
    # Le bandeau footer doit être posé EN DERNIER (z-order top) pour masquer
    # un éventuel débordement de textbox dans le bas du slide.

    initials = cv.get("initials") or initials_from_name(
        cv.get("first_name", ""), cv.get("last_name", "")
    )
    title = cv.get(f"title_{lang}") or cv.get("title", "")
    domain = cv.get(f"domain_{lang}") or cv.get("domain", "")
    _build_header(slide, initials, title, domain)

    payload = _truncate_payload(cv.get(lang, cv), f"{initials} [{lang}]")
    payload = _adaptive_sidebar_trim(payload, f"{initials} [{lang}]")
    payload = _adaptive_main_trim(payload, f"{initials} [{lang}]")
    _check_overflow(payload, f"{initials} [{lang}]")
    _build_sidebar_content(slide, payload, lang)
    _build_main_content(slide, payload, lang)

    # Bandeau masque + texte footer en dernier (z-order top).
    _build_footer_mask(slide)
    _build_footer(slide, lang)


def render(cv_json: dict, out_path: Path) -> Path:
    """Render a single CV to its own .pptx (2 slides FR + EN)."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    logo = _resolve_logo()
    _build_slide(prs, cv_json, "fr", logo)
    _build_slide(prs, cv_json, "en", logo)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path


def render_merged(cv_jsons: list[dict], out_path: Path) -> Path:
    """Render several CVs into ONE combined .pptx (2 slides per CV, ordered
    by input list). Used for batch delivery."""
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    logo = _resolve_logo()
    for cv in cv_jsons:
        _build_slide(prs, cv, "fr", logo)
        _build_slide(prs, cv, "en", logo)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    return out_path


def _slug(s: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9._-]+", "_", s.strip())
    return s.strip("_") or "cv"


def _iter_jsons(path: Path) -> Iterable[Path]:
    return sorted(p for p in path.glob("*.json")) if path.is_dir() else [path]


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Render CV JSONs to Inside Circle PPTX. "
            "Three modes: (a) single JSON + <name>.pptx = 1 CV per file ; "
            "(b) directory + --merged <name>.pptx = ALL CVs in ONE file ; "
            "(c) directory + directory out = one PPT per CV."
        ),
    )
    parser.add_argument("input", type=Path, help="JSON file or directory")
    parser.add_argument("--out", type=Path, default=Path("outputs"))
    parser.add_argument(
        "--merged", action="store_true",
        help="When input is a directory, concatenate every CV into a single "
             ".pptx file (2 slides per CV, ordered alphabetically).",
    )
    args = parser.parse_args()

    if args.merged:
        if not args.input.is_dir():
            raise SystemExit("--merged requires a directory input")
        out_file = args.out if args.out.suffix == ".pptx" \
            else args.out / "all_cvs.pptx"
        jsons = _iter_jsons(args.input)
        cvs = []
        for j in jsons:
            _warnings.clear()
            cv = json.loads(j.read_text(encoding="utf-8"))
            cvs.append(cv)
            if _warnings:
                print(f"  ({len(_warnings)} warnings) {j.name}")
        _warnings.clear()
        rendered = render_merged(cvs, out_file)
        print(f"✓ {len(cvs)} CV(s) → {rendered}")
        return

    out = args.out
    multiple = args.input.is_dir() or out.is_dir() or out.suffix != ".pptx"
    if multiple and out.suffix == ".pptx":
        raise SystemExit("In batch mode, --out must be a directory")

    for j in _iter_jsons(args.input):
        _warnings.clear()
        cv = json.loads(j.read_text(encoding="utf-8"))
        if multiple:
            out_file = (out if out.is_dir() else Path("outputs")) / f"{_slug(j.stem)}.pptx"
        else:
            out_file = out
        rendered = render(cv, out_file)
        suffix = f"  ({len(_warnings)} warnings)" if _warnings else ""
        print(f"✓ {j.name} → {rendered}{suffix}")


if __name__ == "__main__":
    main()
